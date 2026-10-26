# apps/finance/views_reports.py
from collections import defaultdict
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

from dateutil.relativedelta import relativedelta
from django.db.models import Count, Sum
from django.db.models.functions import Coalesce, TruncDate
from django.http import HttpResponse
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.views.decorators.csrf import csrf_exempt
from rest_framework import permissions
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.accounts.models import User
from apps.deals.models import DealRoom
from apps.listings.models import Listing
from apps.transactions.models import Transaction

from .services.user_stats import calculate_user_stats


class IsAdminUser(permissions.BasePermission):
    def has_permission(self, request, view):
        return bool(
            request.user
            and request.user.is_authenticated
            and request.user.is_staff
        )


# ============================================================
# HELPERS
# ============================================================
def _visible_listings_qs():
    return Listing.objects.filter(is_deleted=False).exclude(
        status=Listing.Status.DRAFT,
    )


def _parse_date(s):
    """Parse YYYY-MM-DD → datetime (timezone-aware) au None."""
    if not s:
        return None
    try:
        d = datetime.strptime(s, "%Y-%m-%d").date()
        return timezone.make_aware(
            datetime.combine(d, datetime.min.time()),
            timezone.get_current_timezone(),
        )
    except (ValueError, TypeError):
        return None


def _parse_end_date(s):
    """Parse YYYY-MM-DD → mwisho wa siku (23:59:59) timezone-aware."""
    if not s:
        return None
    try:
        d = datetime.strptime(s, "%Y-%m-%d").date()
        return timezone.make_aware(
            datetime.combine(d, datetime.max.time()),
            timezone.get_current_timezone(),
        )
    except (ValueError, TypeError):
        return None


def _period_range(period, date_from=None, date_to=None):
    """
    Rudisha (start, end) kutoka period au custom date range.
    """
    # ✅ Tumia localtime ili upate muda wa Tanzania (EAT)
    now = timezone.localtime()

    if period == "custom":
        start = _parse_date(date_from)
        end = _parse_end_date(date_to) or now
        if not start:
            return None, None
        return start, end

    if period == "today":
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        return start, now
    if period == "week":
        start = (now - timedelta(days=7))
        return start, now
    if period == "month":
        start = (now - relativedelta(months=1))
        return start, now
    if period == "year":
        start = (now - relativedelta(years=1))
        return start, now
    return None, None  # all


# ============================================================
# REVENUE-BY-MONTH HELPER (revenue halisi, sio GMV)
# ============================================================
def _revenue_for_month(month_start, month_end):
    """
    Hesabu revenue halisi kwa mwezi mmoja kutoka vyanzo vyote.
    """
    from apps.banners.models import BannerAd
    from apps.boosting.models import ListingBoost
    from apps.bundles.models import BundlePurchase
    from apps.leading_fees.models import ListingLeading
    from apps.listings.models import ListingFee
    from apps.transactions.models import Reservation
    from .models import SuccessFeePayment

    total = Decimal("0.00")

    total += ListingFee.objects.filter(
        payment_status="PAID",
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("amount"))["s"] or Decimal("0.00")

    total += ListingBoost.objects.filter(
        payment_status="PAID",
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("amount"))["s"] or Decimal("0.00")

    total += ListingLeading.objects.filter(
        payment_status=ListingLeading.PaymentStatus.PAID,
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("price"))["s"] or Decimal("0.00")

    total += Reservation.objects.filter(
        payment_status="PAID",
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("deposit_amount"))["s"] or Decimal("0.00")

    total += BannerAd.objects.filter(
        payment_status="PAID",
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("amount"))["s"] or Decimal("0.00")

    total += BundlePurchase.objects.filter(
        status=BundlePurchase.Status.PAID,
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("amount"))["s"] or Decimal("0.00")

    total += SuccessFeePayment.objects.filter(
        payment_status="PAID",
        paid_at__gte=month_start, paid_at__lt=month_end,
    ).aggregate(s=Sum("amount"))["s"] or Decimal("0.00")

    return total


# ============================================================
# REPORTS VIEW
# ============================================================
@method_decorator(csrf_exempt, name="dispatch")
class ReportsView(APIView):
    permission_classes = [IsAdminUser]

    def get(self, request):
        # ✅ Tumia localtime
        now = timezone.localtime()

        # ── CUSTOM DATE RANGE (optional) ────────────────────
        period = request.query_params.get("period", "all")
        date_from = request.query_params.get("from")
        date_to = request.query_params.get("to")
        start, end = _period_range(period, date_from, date_to)

        # ── USER STATS ──────────────────────────────────────
        user_stats = calculate_user_stats()

        total_users = user_stats["totalUsers"]
        total_sellers = user_stats["totalSellers"]
        total_buyers = user_stats["totalBuyers"]
        both_roles = user_stats["bothRoles"]
        neither_role = user_stats["neitherRole"]
        total_admins = user_stats["totalAdmins"]

        # ── USERS GROWTH ────────────────────────────────────
        start_day = (now - timedelta(days=29)).date()
        growth_rows = dict(
            User.objects
            .filter(
                is_deleted=False,
                is_staff=False,
                date_joined__date__gte=start_day,
            )
            .annotate(d=TruncDate("date_joined"))
            .values_list("d")
            .annotate(c=Count("id"))
            .values_list("d", "c")
        )
        baseline = User.objects.filter(
            is_deleted=False,
            is_staff=False,
            date_joined__date__lt=start_day,
        ).count()
        users_growth = []
        running = baseline
        for i in range(29, -1, -1):
            day = (now - timedelta(days=i)).date()
            new_count = growth_rows.get(day, 0)
            running += new_count
            users_growth.append({
                "date": day.strftime("%Y-%m-%d"),   # ✅ Tarehe kamili (YYYY-MM-DD)
                "label": day.strftime("%d/%m"),      # ✅ Bado tunatuma label kwa backup
                "count": new_count,                  # ✅ Idadi ya siku hiyo
                "cumulative": running,               # ✅ Jumla kwa ujumla
            })

        # ── LISTINGS ────────────────────────────────────────
        visible_qs = _visible_listings_qs()

        if start and end:
            visible_qs_in_range = visible_qs.filter(
                created_at__gte=start, created_at__lte=end,
            )
        else:
            visible_qs_in_range = visible_qs

        total_listings = visible_qs_in_range.count()
        live_listings = visible_qs_in_range.filter(
            status=Listing.Status.LIVE,
        ).count()
        sold_listings = visible_qs_in_range.filter(
            status=Listing.Status.SOLD,
        ).count()

        listings_growth = []
        for i in range(29, -1, -1):
            day = (now - timedelta(days=i)).date()
            count = visible_qs.filter(created_at__date=day).count()
            listings_growth.append({
                "date": day.strftime("%Y-%m-%d"),  # ✅ Tarehe kamili
                "label": day.strftime("%d/%m"),     # ✅ Label kwa backup
                "count": count,
            })

        # ── DEALS ───────────────────────────────────────────
        deals_qs = DealRoom.objects.all()
        if start and end:
            deals_qs = deals_qs.filter(created_at__gte=start, created_at__lte=end)

        total_deals = deals_qs.count()
        completed_deals = deals_qs.filter(
            status=DealRoom.Status.COMPLETED,
        ).count()

        deals_by_status = dict(
            deals_qs.values_list("status")
            .annotate(c=Count("id"))
            .values_list("status", "c")
        )
        for k in [
            "OPEN", "NEGOTIATING", "ACCEPTED", "RESERVED",
            "AWAITING_FINAL_PAYMENT", "PAYMENT_PROOF_SUBMITTED",
            "COMPLETED", "CANCELLED",
        ]:
            deals_by_status.setdefault(k, 0)

        deals_by_status_frontend = {
            "negotiating": deals_by_status.get("NEGOTIATING", 0),
            "accepted": deals_by_status.get("ACCEPTED", 0),
            "reserved": deals_by_status.get("RESERVED", 0),
            "completed": deals_by_status.get("COMPLETED", 0),
            "cancelled": deals_by_status.get("CANCELLED", 0),
        }

        # ── REVENUE ─────────────────────────────────────────
        revenue_qs = Transaction.objects.filter(
            status=Transaction.Status.COMPLETED,
        )
        if start and end:
            revenue_qs = revenue_qs.filter(
                completed_at__gte=start, completed_at__lte=end,
            )

        total_gmv = (
            revenue_qs.aggregate(s=Sum("agreed_price"))["s"] or 0
        )

        try:
            from .services.revenue import calculate_financial_dashboard
            fees_data = calculate_financial_dashboard("all")
            total_revenue = float(fees_data["total_revenue"])
        except Exception:
            total_revenue = 0.0

        revenue_by_month = []
        _anchor = now.replace(
            day=1, hour=0, minute=0, second=0, microsecond=0,
        )
        for i in range(5, -1, -1):
            month_start = _anchor - relativedelta(months=i)
            month_end = month_start + relativedelta(months=1)
            month_total = _revenue_for_month(month_start, month_end)
            revenue_by_month.append({
                "label": month_start.strftime("%b"),
                "total": float(month_total),
            })

        conversion_rate = (
            (completed_deals / total_deals) * 100 if total_deals else 0
        )

        # ── TOP SELLERS ─────────────────────────────────────
        top_sellers_qs = (
            visible_qs_in_range
            .filter(seller__is_staff=False, seller__is_deleted=False)
            .values("seller__id", "seller__name")
            .annotate(
                views=Coalesce(Sum("views_count"), 0),
                listings=Count("id"),
            )
            .order_by("-views")[:5]
        )
        top_sellers = [
            {
                "name": s["seller__name"] or "—",
                "views": s["views"] or 0,
                "listings": s["listings"],
            }
            for s in top_sellers_qs
        ]

        # ── MOST VIEWED ─────────────────────────────────────
        most_viewed_qs = visible_qs_in_range.order_by("-views_count")[:5]
        most_viewed = [
            {
                "id": l.id,
                "title": l.title,
                "views": l.views_count,
                "price": float(l.price),
            }
            for l in most_viewed_qs
        ]

        # ── TOP CATEGORIES ──────────────────────────────────
        cat_qs = (
            visible_qs_in_range
            .values("category__slug", "category__name")
            .annotate(count=Count("id"), views=Sum("views_count"))
            .order_by("-views")[:5]
        )
        top_categories = [
            {
                "key": c["category__slug"] or "—",
                "name": c["category__name"] or "—",
                "count": c["count"],
                "views": c["views"] or 0,
            }
            for c in cat_qs
        ]

        # ── TOP LOCATIONS ───────────────────────────────────
        loc_qs = (
            visible_qs_in_range
            .values("location")
            .annotate(count=Count("id"), views=Sum("views_count"))
            .order_by("-views")[:5]
        )
        top_locations = [
            {
                "name": l["location"] or "—",
                "count": l["count"],
                "views": l["views"] or 0,
            }
            for l in loc_qs
        ]

        return Response({
            "period": period,
            "from": date_from,
            "to": date_to,
            "totalUsers": total_users,
            "totalSellers": total_sellers,
            "totalBuyers": total_buyers,
            "bothRoles": both_roles,
            "neitherRole": neither_role,
            "totalAdmins": total_admins,
            "usersGrowth": users_growth,

            "totalListings": total_listings,
            "liveListings": live_listings,
            "soldListings": sold_listings,
            "listingsGrowth": listings_growth,

            "totalDeals": total_deals,
            "completedDeals": completed_deals,
            "disputedDeals": 0,
            "dealsByStatus": deals_by_status_frontend,

            "totalGMV": float(total_gmv),
            "totalRevenue": total_revenue,
            "revenueByMonth": revenue_by_month,
            "conversionRate": round(conversion_rate, 2),

            "topSellers": top_sellers,
            "mostViewedListings": most_viewed,
            "topCategories": top_categories,
            "topLocations": top_locations,
        })


# ============================================================
# EXPORT REPORT — CSV / PDF / DOC
# ============================================================
@method_decorator(csrf_exempt, name="dispatch")
class ReportsExportView(APIView):
    permission_classes = [IsAdminUser]

    def get(self, request):
        fmt = (request.query_params.get("format") or "csv").lower()
        period = (request.query_params.get("period") or "month").lower()
        date_from = request.query_params.get("from")
        date_to = request.query_params.get("to")

        if fmt not in ("csv", "pdf", "doc"):
            return Response(
                {"detail": "Format si sahihi. Tumia csv, pdf, au doc."},
                status=400,
            )
        if period not in ("today", "week", "month", "year", "all", "custom"):
            return Response(
                {"detail": "Period si sahihi. Tumia today, week, month, year, all, au custom."},
                status=400,
            )
        if period == "custom" and (not date_from or not date_to):
            return Response(
                {"detail": "Custom period inahitaji 'from' na 'to' (YYYY-MM-DD)."},
                status=400,
            )

        data = ReportsView().get(request).data

        from .services.revenue import get_revenue_records
        period_map = {
            "today": "today",
            "week": "week",
            "month": "month",
            "year": "all",
            "all": "all",
            "custom": "all",
        }
        api_period = period_map.get(period, "all")

        try:
            records = get_revenue_records(period=api_period, source="all")
        except Exception:
            records = []

        if period == "custom":
            start = _parse_date(date_from)
            end = _parse_end_date(date_to)
            if start and end:
                filtered = []
                for r in records:
                    dt = r.get("paid_at") or r.get("created_at")
                    if dt and start <= dt <= end:
                        filtered.append(r)
                records = filtered

        # ✅ Tumia localtime kwa timestamp
        timestamp = timezone.localtime().strftime("%Y%m%d_%H%M")
        suffix = (
            f"{date_from or 'start'}_to_{date_to or 'end'}"
            if period == "custom"
            else period
        )

        if fmt == "csv":
            return self._csv(data, records, suffix, timestamp)
        if fmt == "pdf":
            return self._pdf(data, records, suffix, timestamp)
        return self._doc(data, records, suffix, timestamp)

    # ── CSV ─────────────────────────────────────────────────
    def _csv(self, data, records, period, timestamp):
        import csv
        response = HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = (
            f'attachment; filename="ripoti_{period}_{timestamp}.csv"'
        )
        response.write("\ufeff")
        writer = csv.writer(response)

        writer.writerow(["SokoMkononi — Ripoti ya Mfumo"])
        writer.writerow([f"Kipindi: {period}"])
        # ✅ Tumia localtime
        writer.writerow([f"Tarehe: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}"])
        writer.writerow([])

        writer.writerow(["=== WATUMIAJI ==="])
        writer.writerow(["Jumla ya Watumiaji", data.get("totalUsers", 0)])
        writer.writerow(["Wauzaji Pekee", data.get("totalSellers", 0)])
        writer.writerow(["Wanunuzi Pekee", data.get("totalBuyers", 0)])
        writer.writerow(["Wote Wawili", data.get("bothRoles", 0)])
        writer.writerow(["Wapya", data.get("neitherRole", 0)])
        writer.writerow(["Admins", data.get("totalAdmins", 0)])
        writer.writerow([])

        writer.writerow(["=== MALI ==="])
        writer.writerow(["Jumla ya Mali", data.get("totalListings", 0)])
        writer.writerow(["Mali Hai", data.get("liveListings", 0)])
        writer.writerow(["Zimeuzwa", data.get("soldListings", 0)])
        writer.writerow([])

        writer.writerow(["=== DEALS ==="])
        writer.writerow(["Jumla ya Deals", data.get("totalDeals", 0)])
        writer.writerow(["Zimekamilika", data.get("completedDeals", 0)])
        for k, v in (data.get("dealsByStatus") or {}).items():
            writer.writerow([f"  {k}", v])
        writer.writerow([])

        writer.writerow(["=== MAPATO ==="])
        writer.writerow([
            "Jumla ya Mapato",
            f"TZS {float(data.get('totalRevenue', 0)):,.0f}",
        ])
        writer.writerow([
            "Jumla ya GMV",
            f"TZS {float(data.get('totalGMV', 0)):,.0f}",
        ])
        writer.writerow([
            "Conversion Rate",
            f"{float(data.get('conversionRate', 0)):.2f}%",
        ])
        writer.writerow([])

        writer.writerow(["=== WAUZAJI BORA ==="])
        writer.writerow(["#", "Jina", "Mali", "Views"])
        for i, s in enumerate(data.get("topSellers", []), 1):
            writer.writerow([i, s.get("name"), s.get("listings", 0), s.get("views", 0)])
        writer.writerow([])

        writer.writerow(["=== MALI ZINAZOONEKANA ZAIDI ==="])
        writer.writerow(["#", "Jina", "Views", "Bei"])
        for i, l in enumerate(data.get("mostViewedListings", []), 1):
            writer.writerow([
                i, l.get("title"), l.get("views", 0),
                f"TZS {float(l.get('price', 0)):,.0f}",
            ])
        writer.writerow([])

        writer.writerow(["=== KATEGORIA ZINAZOONEKANA ZAIDI ==="])
        writer.writerow(["#", "Kategoria", "Mali", "Views"])
        for i, c in enumerate(data.get("topCategories", []), 1):
            writer.writerow([
                i, c.get("name") or c.get("key"), c.get("count", 0), c.get("views", 0),
            ])
        writer.writerow([])

        writer.writerow(["=== MAENEO YANAYOONEKANA ZAIDI ==="])
        writer.writerow(["#", "Eneo", "Mali", "Views"])
        for i, l in enumerate(data.get("topLocations", []), 1):
            writer.writerow([i, l.get("name"), l.get("count", 0), l.get("views", 0)])
        writer.writerow([])

        writer.writerow(["=== MIAMALA ILIYOKAMILIKA ==="])
        writer.writerow(["Tarehe", "Kumbukumbu", "Aina", "Kichwa", "Kiasi (TZS)", "Hali"])
        for r in records:
            date_val = r.get("paid_at") or r.get("created_at")
            date_str = date_val.strftime("%Y-%m-%d") if date_val else ""
            writer.writerow([
                date_str,
                str(r.get("payment_reference") or "—")[:20],
                r.get("source_label") or "—",
                (r.get("listing_title") or "")[:40],
                f"{float(r.get('amount', 0)):,.0f}",
                r.get("payment_status") or "—",
            ])

        return response

    # ── PDF ─────────────────────────────────────────────────
    def _pdf(self, data, records, period, timestamp):
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.platypus import (
                Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
            )
        except ImportError:
            return Response(
                {"detail": "reportlab haipo. Wasiliana na admin."},
                status=500,
            )

        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer, pagesize=A4,
            rightMargin=1.5 * cm, leftMargin=1.5 * cm,
            topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        )
        styles = getSampleStyleSheet()
        elements = []

        elements.append(Paragraph(
            "<b>SokoMkononi — Ripoti ya Mfumo</b>", styles["Title"],
        ))
        elements.append(Spacer(1, 0.3 * cm))
        # ✅ Tumia localtime
        elements.append(Paragraph(
            f"Kipindi: <b>{period}</b> · Tarehe: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}",
            styles["Normal"],
        ))
        elements.append(Spacer(1, 0.5 * cm))

        def make_table(rows, headers):
            tbl_data = [headers] + rows
            t = Table(tbl_data, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#101A2E")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E6E2D6")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F3EC")]),
            ]))
            return t

        elements.append(Paragraph("<b>1. Watumiaji</b>", styles["Heading2"]))
        elements.append(make_table([
            ["Jumla ya Watumiaji", str(data.get("totalUsers", 0))],
            ["Wauzaji Pekee", str(data.get("totalSellers", 0))],
            ["Wanunuzi Pekee", str(data.get("totalBuyers", 0))],
            ["Wote Wawili", str(data.get("bothRoles", 0))],
            ["Wapya", str(data.get("neitherRole", 0))],
            ["Admins", str(data.get("totalAdmins", 0))],
        ], ["Kipimo", "Thamani"]))
        elements.append(Spacer(1, 0.5 * cm))

        elements.append(Paragraph("<b>2. Mali</b>", styles["Heading2"]))
        elements.append(make_table([
            ["Jumla ya Mali", str(data.get("totalListings", 0))],
            ["Mali Hai", str(data.get("liveListings", 0))],
            ["Zimeuzwa", str(data.get("soldListings", 0))],
        ], ["Kipimo", "Thamani"]))
        elements.append(Spacer(1, 0.5 * cm))

        elements.append(Paragraph("<b>3. Deals</b>", styles["Heading2"]))
        elements.append(make_table([
            ["Jumla ya Deals", str(data.get("totalDeals", 0))],
            ["Zimekamilika", str(data.get("completedDeals", 0))],
        ], ["Kipimo", "Thamani"]))
        elements.append(Spacer(1, 0.5 * cm))

        elements.append(Paragraph("<b>4. Mapato</b>", styles["Heading2"]))
        elements.append(make_table([
            ["Jumla ya Mapato", f"TZS {float(data.get('totalRevenue', 0)):,.0f}"],
            ["Jumla ya GMV", f"TZS {float(data.get('totalGMV', 0)):,.0f}"],
            ["Conversion Rate", f"{float(data.get('conversionRate', 0)):.2f}%"],
        ], ["Kipimo", "Thamani"]))

        doc.build(elements)
        buffer.seek(0)
        response = HttpResponse(buffer.read(), content_type="application/pdf")
        response["Content-Disposition"] = (
            f'attachment; filename="ripoti_{period}_{timestamp}.pdf"'
        )
        return response

    # ── DOC ─────────────────────────────────────────────────
    def _doc(self, data, records, period, timestamp):
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return Response(
                {"detail": "python-docx haipo. Wasiliana na admin."},
                status=500,
            )

        doc = Document()
        h = doc.add_heading("SokoMkononi — Ripoti ya Mfumo", 0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # ✅ Tumia localtime
        sub = doc.add_paragraph(
            f"Kipimo: {period} · Tarehe: {timezone.localtime().strftime('%Y-%m-%d %H:%M')}"
        )
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

        def add_section(title, rows, headers):
            doc.add_heading(title, level=1)
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = "Light Grid Accent 1"
            hdr = t.rows[0].cells
            for i, h in enumerate(headers):
                hdr[i].text = h
            for r in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = str(v)

        add_section("1. Watumiaji", [
            ["Jumla ya Watumiaji", data.get("totalUsers", 0)],
            ["Wauzaji Pekee", data.get("totalSellers", 0)],
            ["Wanunuzi Pekee", data.get("totalBuyers", 0)],
            ["Wote Wawili", data.get("bothRoles", 0)],
            ["Wapya", data.get("neitherRole", 0)],
            ["Admins", data.get("totalAdmins", 0)],
        ], ["Kipimo", "Thamani"])

        add_section("2. Mali", [
            ["Jumla ya Mali", data.get("totalListings", 0)],
            ["Mali Hai", data.get("liveListings", 0)],
            ["Zimeuzwa", data.get("soldListings", 0)],
        ], ["Kipimo", "Thamani"])

        add_section("3. Deals", [
            ["Jumla ya Deals", data.get("totalDeals", 0)],
            ["Zimekamilika", data.get("completedDeals", 0)],
        ], ["Kipimo", "Thamani"])

        add_section("4. Mapato", [
            ["Jumla ya Mapato", f"TZS {float(data.get('totalRevenue', 0)):,.0f}"],
            ["Jumla ya GMV", f"TZS {float(data.get('totalGMV', 0)):,.0f}"],
            ["Conversion Rate", f"{float(data.get('conversionRate', 0)):.2f}%"],
        ], ["Kipimo", "Thamani"])

        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        response = HttpResponse(
            buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response["Content-Disposition"] = (
            f'attachment; filename="ripoti_{period}_{timestamp}.docx"'
        )
        return response